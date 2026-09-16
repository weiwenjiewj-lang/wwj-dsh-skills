#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_docx.py —— 把交付稿 md 输出为方案报告书 Word 文件

严格照 `references/table-format.md`（＝办水保函〔2026〕232号-附件2 原文）排版：
  正文 小四号仿宋；数字与英文 小四号 Times New Roman
  页眉＝相应章节名称，五号仿宋；页脚＝编制单位名称+页码，五号仿宋
  封面 湖蓝色；项目名称 加粗二号宋体；"水土保持方案报告书" 小初号黑体
  扉页 版式同封面；责任页 报告书名称三号黑体，其余四号（标签黑体、姓名分工宋体）
  目录 两级；标题三号黑体，其他四号仿宋
  表题 表上方居中 五号黑体；表注 表下方 小五号宋体；表内文字 五号仿宋

本脚本相对"临时拼装脚本"修正的四个硬伤：
  1) 封面**只生成一次**（封面由本脚本统一产出，md 前置部分不再另出封面）；
  2) **逐章分节**，页眉随章变化（此前全书一个分节 → 页眉无法逐章不同）；
  3) 封面/扉页/责任页**所在节不含页眉页脚**；
  4) 表头行**跨页重复**且用**黑体**加粗。

用法
----
    python build_docx.py --md 交付稿.md --out 方案.docx \
        --project "XX项目" --owner "建设单位" --compiler "编制单位" \
        --date "2026 年 9 月" [--no-open] [--toc-level 2]

退出码：0 成功 · 2 输入/用法错误
"""

import sys
sys.dont_write_bytecode = True

import argparse
import os
import re
import subprocess

import docx
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

# ---------------------------------------------------------------- 规格常量
CN_BODY = '仿宋'
CN_HEI = '黑体'
CN_SONG = '宋体'
EN = 'Times New Roman'

FS_BODY = Pt(12)         # 小四
FS_TBL = Pt(10.5)        # 五号
FS_HDR = Pt(10.5)        # 页眉页脚 五号
FS_CAP = Pt(10.5)        # 表题 五号黑体
FS_NOTE = Pt(9)          # 表注 小五宋体
FS_H1 = Pt(16)           # 章标题 三号
FS_H2 = Pt(14)           # 节标题 四号
FS_H3 = Pt(12)           # 小节标题 小四
FS_COVER_PROJ = Pt(22)   # 二号
FS_COVER_TITLE = Pt(36)  # 小初
FS_COVER_UNIT = Pt(16)   # 三号宋体

HULAN = RGBColor(0x1F, 0x6F, 0x9C)   # 湖蓝色（附件2 §二）

ROLES = ['批　准', '核　定', '审　查', '校　核', '项目负责人', '编　写']

RE_H1 = re.compile(r'^#\s+(.*)$')
RE_H2 = re.compile(r'^##\s+(.*)$')
RE_H3 = re.compile(r'^###\s+(.*)$')
RE_H4 = re.compile(r'^####\s+(.*)$')
RE_CH = re.compile(r'^第\s*(\d+)\s*章[\s　]*(.*)$')
RE_CAP = re.compile(r'^表\s*\d+([-—－.]\d+)?\s*[　\s]\s*\S')
RE_TBL_LINE = re.compile(r'^\s*\|.*\|\s*$')
RE_SEP = re.compile(r'^\s*\|[\s:|-]+\|\s*$')
RE_NOTE = re.compile(r'^\s*注\s*[：:1-9０-９]')


def set_run(run, cn=CN_BODY, size=FS_BODY, bold=False, color=None):
    """西文/数字 Times New Roman，中文 cn —— 附件2 §六.1 要求中英文分别设字体"""
    run.font.name = EN
    run.font.size = size
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rPr.insert(0, rf)
    rf.set(qn('w:ascii'), EN)
    rf.set(qn('w:hAnsi'), EN)
    rf.set(qn('w:eastAsia'), cn)


def add_para(doc, text='', cn=CN_BODY, size=FS_BODY, bold=False, align=None,
             indent=True, space_before=0, space_after=0, style=None, color=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.5
    if indent and align is None:
        pf.first_line_indent = Pt(24)   # 首行缩进 2 字符
    if text:
        r = p.add_run(text)
        set_run(r, cn=cn, size=size, bold=bold, color=color)
    return p


def add_field(paragraph, instr, placeholder='1'):
    """插入 Word 域（页码 / 目录）"""
    r = paragraph.add_run()
    fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr
    fs = OxmlElement('w:fldChar'); fs.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t'); t.text = placeholder
    fe = OxmlElement('w:fldChar'); fe.set(qn('w:fldCharType'), 'end')
    r._element.append(fc); r._element.append(it); r._element.append(fs)
    r._element.append(t); r._element.append(fe)
    return r


def style_hf(section, header_text, footer_unit, with_page=True):
    """页眉＝章节名称（五号仿宋）；页脚＝编制单位名称+页码（五号仿宋）"""
    hdr = section.header
    hdr.is_linked_to_previous = False
    hp = hdr.paragraphs[0] if hdr.paragraphs else hdr.add_paragraph()
    hp.text = ''
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if header_text:
        set_run(hp.add_run(header_text), cn=CN_BODY, size=FS_HDR)

    ftr = section.footer
    ftr.is_linked_to_previous = False
    fp = ftr.paragraphs[0] if ftr.paragraphs else ftr.add_paragraph()
    fp.text = ''
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if footer_unit:
        set_run(fp.add_run(footer_unit), cn=CN_BODY, size=FS_HDR)
    if with_page:
        set_run(fp.add_run('　'), cn=CN_BODY, size=FS_HDR)
        set_run(fp.add_run('第 '), cn=CN_BODY, size=FS_HDR)
        r = add_field(fp, ' PAGE ')
        set_run(r, cn=CN_BODY, size=FS_HDR)
        set_run(fp.add_run(' 页'), cn=CN_BODY, size=FS_HDR)


def set_repeat_header(table):
    """表头行跨页重复"""
    if not table.rows:
        return
    tr = table.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    if trPr.find(qn('w:tblHeader')) is None:
        el = OxmlElement('w:tblHeader')
        el.set(qn('w:val'), 'true')
        trPr.append(el)


def style_table(table, header_hei=True):
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                if i == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    if i == 0 and header_hei:
                        set_run(r, cn=CN_HEI, size=FS_TBL, bold=True)
                    else:
                        set_run(r, cn=CN_BODY, size=FS_TBL)
    set_repeat_header(table)


def parse_table(lines, i):
    """从 lines[i] 起解析 markdown 表格，返回 (rows, next_i)"""
    rows = []
    while i < len(lines) and RE_TBL_LINE.match(lines[i]):
        if not RE_SEP.match(lines[i]):
            cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
            rows.append(cells)
        i += 1
    return rows, i


def strip_md(t):
    t = re.sub(r'\*\*(.+?)\*\*', r'\1', t)
    t = re.sub(r'(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)', r'\1', t)
    t = re.sub(r'`(.+?)`', r'\1', t)
    t = re.sub(r'^\s*[-*+]\s+', '', t)
    t = re.sub(r'^\s*\d+[.)]\s+', '', t)
    return t.strip()


def build(md_path, out_path, project, owner, compiler, date_text,
          report_name='水土保持方案报告书', toc_level=2, responsibility=None):
    lines = open(md_path, encoding='utf-8').read().split('\n')
    doc = Document()

    st = doc.styles['Normal']
    st.font.name = EN
    st.font.size = FS_BODY
    st.element.rPr.rFonts.set(qn('w:eastAsia'), CN_BODY)
    st.paragraph_format.line_spacing = 1.5

    sec0 = doc.sections[0]
    sec0.page_height, sec0.page_width = Cm(29.7), Cm(21.0)
    sec0.top_margin = sec0.bottom_margin = Cm(2.54)
    sec0.left_margin = sec0.right_margin = Cm(3.17)
    # 封面/扉页/责任页：无页眉页脚
    style_hf(sec0, '', '', with_page=False)

    center = WD_ALIGN_PARAGRAPH.CENTER
    # ---------------- 封面（唯一一次）----------------
    for _ in range(3):
        add_para(doc, '', indent=False)
    add_para(doc, project, cn=CN_SONG, size=FS_COVER_PROJ, bold=True,
             align=center, indent=False, space_after=10, color=HULAN)
    add_para(doc, report_name, cn=CN_HEI, size=FS_COVER_TITLE, bold=False,
             align=center, indent=False, space_after=30, color=HULAN)
    for _ in range(6):
        add_para(doc, '', indent=False)
    add_para(doc, '建设单位：' + owner, cn=CN_SONG, size=FS_COVER_UNIT,
             align=center, indent=False, space_after=6)
    add_para(doc, '编制单位：' + compiler, cn=CN_SONG, size=FS_COVER_UNIT,
             align=center, indent=False, space_after=6)
    add_para(doc, date_text, cn=CN_SONG, size=FS_COVER_UNIT,
             align=center, indent=False)

    # ---------------- 扉页（版式同封面，供盖章）----------------
    doc.add_page_break()
    for _ in range(3):
        add_para(doc, '', indent=False)
    add_para(doc, project, cn=CN_SONG, size=FS_COVER_PROJ, bold=True,
             align=center, indent=False, space_after=10, color=HULAN)
    add_para(doc, report_name, cn=CN_HEI, size=FS_COVER_TITLE, bold=False,
             align=center, indent=False, space_after=30, color=HULAN)
    for _ in range(6):
        add_para(doc, '', indent=False)
    add_para(doc, '建设单位：' + owner + '　（盖章）', cn=CN_SONG, size=FS_COVER_UNIT,
             align=center, indent=False, space_after=6)
    add_para(doc, '编制单位：' + compiler + '　（盖章）', cn=CN_SONG, size=FS_COVER_UNIT,
             align=center, indent=False, space_after=6)
    add_para(doc, date_text, cn=CN_SONG, size=FS_COVER_UNIT,
             align=center, indent=False)

    # ---------------- 责任页 ----------------
    doc.add_page_break()
    add_para(doc, report_name + '责任页', cn=CN_HEI, size=FS_H1, bold=True,
             align=center, indent=False, space_after=12)
    rows = responsibility or [[r, '【待填：姓名】', '【待填：职务/职称】', '【待填：分工】', ''] for r in ROLES]
    t = doc.add_table(rows=1 + len(rows), cols=5)
    hdr = ['职　责', '姓　名', '职务/职称', '分　工', '签　名']
    for j, h in enumerate(hdr):
        t.rows[0].cells[j].text = h
    for i, r in enumerate(rows, 1):
        for j in range(5):
            t.rows[i].cells[j].text = r[j] if j < len(r) else ''
    style_table(t, header_hei=True)
    # 标签列黑体、其余宋体（附件2 §四）
    for i, row in enumerate(t.rows):
        for p in row.cells[0].paragraphs:
            for r in p.runs:
                set_run(r, cn=CN_HEI, size=Pt(14) if i else FS_TBL, bold=(i > 0))
        for j in range(1, 5):
            for p in row.cells[j].paragraphs:
                for r in p.runs:
                    set_run(r, cn=CN_SONG, size=Pt(14) if i else FS_TBL, bold=False)
    add_para(doc, '注：本表须加盖建设单位与编制单位公章，编制人员须亲笔签名。',
             cn=CN_SONG, size=FS_NOTE, indent=False, space_before=6)

    # ---------------- 目录（独立分节；真目录域）----------------
    sec_toc = doc.add_section(WD_SECTION.NEW_PAGE)
    style_hf(sec_toc, '目　录', compiler)
    add_para(doc, '目　　录', cn=CN_HEI, size=FS_H1, bold=True,
             align=center, indent=False, space_after=12)
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    r = add_field(p, ' TOC \\o "1-%d" \\h \\z \\u ' % toc_level,
                  placeholder='（在 Word 中按 F9 或「引用→更新目录」生成页码）')
    set_run(r, cn=CN_BODY, size=Pt(14))

    # ---------------- 正文：逐章分节 ----------------
    cur_sec = None
    chapter_no = 0
    pending_caption = None
    prev_was_table = False

    def ensure_section(title):
        """开启新分节，页眉＝章节名称"""
        nonlocal cur_sec
        cur_sec = doc.add_section(WD_SECTION.NEW_PAGE)
        style_hf(cur_sec, title, compiler)
        return cur_sec

    i = 0
    while i < len(lines):
        raw = lines[i]
        s = raw.strip()

        if not s:
            i += 1
            prev_was_table = False
            continue
        if re.match(r'^-{3,}$', s) or re.match(r'^\*{3,}$', s):
            i += 1
            continue

        # 表格
        if RE_TBL_LINE.match(raw):
            rows, i2 = parse_table(lines, i)
            if rows:
                if pending_caption:
                    add_para(doc, pending_caption, cn=CN_HEI, size=FS_CAP,
                             align=center, indent=False, space_before=6, space_after=2)
                    pending_caption = None
                ncol = max(len(r) for r in rows)
                tb = doc.add_table(rows=len(rows), cols=ncol)
                for ri, r in enumerate(rows):
                    for ci in range(ncol):
                        tb.rows[ri].cells[ci].text = r[ci] if ci < len(r) else ''
                style_table(tb)
                prev_was_table = True
            i = i2
            continue

        # 表题（表上方居中 五号黑体）
        if RE_CAP.match(s) and not prev_was_table:
            pending_caption = strip_md(s)
            i += 1
            continue

        # 表注（表下方 小五号宋体）
        if prev_was_table and RE_NOTE.match(s):
            add_para(doc, strip_md(s), cn=CN_SONG, size=FS_NOTE,
                     indent=False, space_after=6)
            prev_was_table = False
            i += 1
            continue
        prev_was_table = False

        # 一级：章
        m = RE_H1.match(s)
        if m:
            title = strip_md(m.group(1))
            mc = RE_CH.match(title)
            if mc:
                chapter_no += 1
                ensure_section(title)
                add_para(doc, title, cn=CN_HEI, size=FS_H1, bold=True,
                         align=center, indent=False, space_before=0, space_after=14,
                         style='Heading 1')
            else:
                # 非章的一级标题（附表/附件/附图/附录等）→ 各自分节
                ensure_section(title)
                add_para(doc, title, cn=CN_HEI, size=FS_H1, bold=True,
                         align=center, indent=False, space_before=0, space_after=14,
                         style='Heading 1')
            i += 1
            continue

        m = RE_H2.match(s)
        if m:
            title = strip_md(m.group(1))
            add_para(doc, title, cn=CN_HEI, size=FS_H2, bold=True,
                     indent=False, space_before=10, space_after=6, style='Heading 2')
            i += 1
            continue

        m = RE_H3.match(s)
        if m:
            title = strip_md(m.group(1))
            add_para(doc, title, cn=CN_HEI, size=FS_H3, bold=True,
                     indent=False, space_before=8, space_after=4)
            i += 1
            continue

        m = RE_H4.match(s)
        if m:
            title = strip_md(m.group(1))
            add_para(doc, title, cn=CN_HEI, size=FS_H3, bold=True,
                     indent=False, space_before=6, space_after=4)
            i += 1
            continue

        # 普通段落
        txt = strip_md(s)
        if not txt:
            i += 1
            continue
        add_para(doc, txt, cn=CN_BODY, size=FS_BODY, indent=True)
        i += 1

    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or '.', exist_ok=True)
    doc.save(out_path)
    return doc


def main():
    ap = argparse.ArgumentParser(description='交付稿 md → 方案报告书 Word')
    ap.add_argument('--md', required=True, help='finalize.py 产出的交付稿 md')
    ap.add_argument('--out', required=True, help='输出 .docx（命名须带日期）')
    ap.add_argument('--project', required=True)
    ap.add_argument('--owner', default='【待填：建设单位】')
    ap.add_argument('--compiler', default='【待填：编制单位】')
    ap.add_argument('--date', dest='date_text', default='')
    ap.add_argument('--toc-level', type=int, default=2)
    ap.add_argument('--no-open', action='store_true', help='不自动打开')
    args = ap.parse_args()

    if not os.path.exists(args.md):
        print('输入错误：找不到 %s' % args.md, file=sys.stderr)
        return 2
    if not args.out.lower().endswith('.docx'):
        print('输入错误：--out 须为 .docx', file=sys.stderr)
        return 2
    if _inside_skill(args.out):
        print('拒绝写入技能包/知识库：%s' % args.out, file=sys.stderr)
        return 2
    if not args.date_text:
        import datetime
        d = datetime.date.today()
        args.date_text = '%d 年 %d 月' % (d.year, d.month)

    doc = build(args.md, args.out, args.project, args.owner, args.compiler,
                args.date_text, toc_level=args.toc_level)

    sys.stdout.reconfigure(encoding='utf-8')
    print('Word 文件已生成：%s' % os.path.abspath(args.out))
    print('  段落 %d ; 表格 %d ; 分节 %d'
          % (len(doc.paragraphs), len(doc.tables), len(doc.sections)))
    print('  页眉：逐章分节，页眉＝章名；封面/扉页/责任页所在节无页眉页脚')
    print('  目录：两级 TOC 域（在 Word 中更新域即出页码）')

    if not args.no_open:
        try:
            if sys.platform.startswith('win'):
                os.startfile(os.path.abspath(args.out))
            elif sys.platform == 'darwin':
                subprocess.Popen(['open', os.path.abspath(args.out)])
            else:
                subprocess.Popen(['xdg-open', os.path.abspath(args.out)])
            print('  已自动打开')
        except Exception as e:
            print('  自动打开失败（请手动打开）：%s' % e)
    return 0


def _inside_skill(path):
    try:
        ap = os.path.abspath(path)
    except Exception:
        return False
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    for root in (here, os.environ.get('DSH_WS_VAULT', '')):
        if root and ap.lower().startswith(os.path.abspath(root).lower()):
            return True
    if os.sep + 'vault' + os.sep in ap.lower():
        return True
    return False


if __name__ == '__main__':
    sys.exit(main())
